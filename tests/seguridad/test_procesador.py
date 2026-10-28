"""Procesamiento de archivos subidos: cada tipo legítimo se convierte y los maliciosos fallan de forma controlada.

El worker corre con bin/worker-pruebas.sh: mismo aislamiento que el servicio estudio-worker (requiere root).
"""
from __future__ import annotations

import io
import os
import subprocess
import wave
import zipfile
from pathlib import Path

import pytest

from conftest import RAIZ, sql

pytestmark = pytest.mark.skipif(os.geteuid() != 0, reason="el worker aislado se lanza con systemd-run (root)")


def correr_worker():
    r = subprocess.run([str(RAIZ / "bin/worker-pruebas.sh")], capture_output=True, text=True, timeout=900)
    assert r.returncode == 0, r.stderr[-2000:]


# ───── Archivos de ejemplo ─────

def pdf() -> bytes:
    import pymupdf
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Tema 1. Cinemática\nLa velocidad es el cambio de posición.", fontsize=12)
    return doc.tobytes()


def docx() -> bytes:
    import docx as d
    doc = d.Document()
    doc.add_heading("Apuntes de Química", 1)
    doc.add_paragraph("El agua es H2O.")
    b = io.BytesIO(); doc.save(b); return b.getvalue()


def pptx() -> bytes:
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Leyes de Newton"
    s.placeholders[1].text = "F = m a"
    b = io.BytesIO(); p.save(b); return b.getvalue()


def xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Materia", "Calificación"]); wb.active.append(["Física", 9])
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def png(lado=64) -> bytes:
    from PIL import Image
    b = io.BytesIO(); Image.new("RGB", (lado, lado), (200, 30, 30)).save(b, "PNG"); return b.getvalue()


def wav() -> bytes:
    b = io.BytesIO()
    with wave.open(b, "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(16000)
        w.writeframes(b"\x00\x00" * 16000 * 2)
    return b.getvalue()


def mp4(tmp: Path) -> bytes:
    salida = tmp / "v.mp4"
    subprocess.run(["ffmpeg", "-v", "error", "-f", "lavfi", "-i", "testsrc=duration=2:size=160x120:rate=10", "-pix_fmt", "yuv420p", str(salida)], check=True)
    return salida.read_bytes()


def bomba_docx() -> bytes:
    """Un .docx de ~600 KB que se infla a 600 MB."""
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        with z.open("word/document.xml", "w", force_zip64=True) as f:
            bloque = b"\x00" * (1 << 20)
            for _ in range(600):
                f.write(bloque)
    return b.getvalue()


def bomba_png() -> bytes:
    """PNG pequeño de 10 000 × 10 000 píxeles (100 MP)."""
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = None
    b = io.BytesIO(); Image.new("1", (10000, 10000)).save(b, "PNG", optimize=True); return b.getvalue()


def subir(c, nombre: str, contenido: bytes, mime: str, seccion="apuntes"):
    return c.s.post(c.url("api/subir.php"), headers={"X-CSRF": c.csrf}, data={"seccion": seccion, "titulo": nombre},
                    files={"archivo": (nombre, contenido, mime)})


# ───── Pruebas ─────

def test_procesa_cada_tipo_y_rechaza_maliciosos(como, nuevo_usuario, tmp_path):
    c = como(nuevo_usuario("profesor"))
    casos = {
        "apuntes.pdf": (pdf(), "application/pdf", "listo"),
        "quimica.docx": (docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "listo"),
        "newton.pptx": (pptx(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "listo"),
        "notas.xlsx": (xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "listo"),
        "foto.png": (png(), "image/png", "listo"),
        "audio.wav": (wav(), "audio/wav", "listo"),
        "video.mp4": (mp4(tmp_path), "video/mp4", "listo"),
        "texto.txt": ("hola <script>alert(1)</script>".encode(), "text/plain", "listo"),
        "roto.pdf": (b"%PDF-1.7\n" + os.urandom(3000), "application/pdf", "error"),
        "bomba.docx": (bomba_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "error"),
        "bomba.png": (bomba_png(), "image/png", "error"),
    }
    ids = {}
    for nombre, (contenido, mime, _) in casos.items():
        r = subir(c, nombre, contenido, mime)
        assert r.status_code == 201, (nombre, r.text)
        ids[nombre] = r.json()["archivo"]["id"]
    correr_worker()
    estados = {f["titulo"]: f for f in sql("SELECT titulo, estado, error FROM archivos")}
    for nombre, (_, _, esperado) in casos.items():
        assert estados[nombre]["estado"] == esperado, (nombre, estados[nombre])
    assert "sospechosa" in estados["bomba.docx"]["error"]
    # La vista del .txt muestra el script como texto, nunca como etiqueta.
    vista = c.get("api/archivo.php", id=ids["texto.txt"], modo="vista").text
    assert "<script>" not in vista and "&lt;script&gt;" in vista


def test_rechaza_contenido_que_no_coincide_con_la_extension(como, nuevo_usuario):
    c = como(nuevo_usuario("alumno"))
    php = b"<?php system($_GET['c']); ?>"
    assert subir(c, "tarea.pdf", php, "application/pdf").status_code == 415
    assert subir(c, "shell.php", php, "application/x-php").status_code == 415
    assert subir(c, "imagen.svg", b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>', "image/svg+xml").status_code == 415
    assert subir(c, "vacio.txt", b"", "text/plain").status_code in (400, 422)


def test_nombre_con_ruta_no_escapa_de_storage(como, nuevo_usuario):
    c = como(nuevo_usuario("alumno"))
    r = subir(c, "../../../../tmp/escape.txt", b"hola", "text/plain")
    assert r.status_code == 201
    assert not Path("/tmp/escape.txt").exists()
    assert r.json()["archivo"]["nombre_original"] == "escape.txt"
